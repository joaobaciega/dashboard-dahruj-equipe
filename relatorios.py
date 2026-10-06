"""Relatórios de desempenho em Excel."""
import datetime as dt
import io

import pandas as pd
from indicadores import agg_by, mlabel

def gerar_excel_ranking(f, periodo, filtros_txt):
    """Gera, em memória, um relatório Excel formatado do Ranking da seleção atual.
    Aba 1: ranking por consultor no trimestre. Aba 2: resumo mensal."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    NAVY, WHITE = "FF1F3864", "FFFFFFFF"
    thin = Side(style="thin", color="FFD9D9D9")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _v(x):  # NaN -> None (célula vazia em vez de "nan")
        return None if (isinstance(x, float) and pd.isna(x)) else x

    def _i(x):
        return int(x) if pd.notna(x) else 0

    def cab(ws, headers, row):
        for j, h in enumerate(headers, 1):
            c = ws.cell(row, j, h)
            c.font = Font(name="Arial", bold=True, color=WHITE, size=10)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BD

    wb = Workbook()

    # ---- Aba 1: Ranking por consultor (trimestre) ----
    ws = wb.active
    ws.title = "Ranking Trimestre"
    ws.cell(1, 1, "Ranking Vendas Refis para Palhetas — Dahruj").font = \
        Font(name="Arial", bold=True, size=14, color=NAVY)
    for r, txt in ((2, f"Período: {periodo}"), (3, f"Filtros: {filtros_txt}"),
                   (4, f"Gerado em: {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}")):
        ws.cell(r, 1, txt).font = Font(name="Arial", size=9, color="FF666666")

    rk = agg_by(f, ["consultor", "unidade", "marca"]).sort_values("total_geral", ascending=False)
    overall = agg_by(f.assign(_g=1), "_g").iloc[0]
    head = ["#", "Consultor", "Unidade", "Marca", "Passagens", "Refil Diant.",
            "Aprov. Diant.", "Refil Tras.", "Aprov. Tras.", "Faturamento (R$)"]
    H = 6
    cab(ws, head, H)
    for i, (_, row) in enumerate(rk.iterrows()):
        r = H + 1 + i
        vals = [i + 1, row["consultor"], row["unidade"], row["marca"],
                row["passagens"], row["refil_diant"], row["aproveitamento"],
                row["refil_tras"], row["aproveitamento_tras"], row["total_geral"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, _v(v)); c.font = Font(name="Arial", size=10); c.border = BD
            if j in (1, 5, 6, 7, 8, 9):
                c.alignment = Alignment(horizontal="center")
            if j in (5, 6, 8):
                c.number_format = "#,##0"
            elif j in (7, 9):
                c.number_format = "0.0%"
            elif j == 10:
                c.number_format = "R$ #,##0.00"
        if i % 2 == 1:
            for j in range(1, len(head) + 1):
                ws.cell(r, j).fill = PatternFill("solid", fgColor="FFF4F6FA")
    tr = H + 1 + len(rk)
    ws.cell(tr, 4, "TOTAL").font = Font(name="Arial", bold=True, size=10)
    for j, v, fmt in ((5, _i(overall["passagens"]), "#,##0"),
                      (6, _i(overall["refil_diant"]), "#,##0"),
                      (7, _v(overall["aproveitamento"]), "0.0%"),
                      (8, _i(overall["refil_tras"]), "#,##0"),
                      (9, _v(overall["aproveitamento_tras"]), "0.0%"),
                      (10, _v(overall["total_geral"]), "R$ #,##0.00")):
        c = ws.cell(tr, j, v); c.font = Font(name="Arial", bold=True, size=10)
        c.border = BD; c.number_format = fmt
        if j in (5, 6, 7, 8, 9):
            c.alignment = Alignment(horizontal="center")
    for j, w in enumerate([5, 28, 16, 9, 11, 12, 12, 11, 12, 15], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"A{H + 1}"

    # ---- Aba 2: Resumo mensal ----
    ws2 = wb.create_sheet("Resumo Mensal")
    mensal = agg_by(f, "mes").sort_values("mes")
    mensal["lbl"] = mensal["mes"].apply(mlabel)
    head2 = ["Mês", "Passagens", "Refil Diant.", "Aprov. Diant.", "Refil Tras.",
             "Aprov. Tras.", "Fat. Diant. (R$)", "Fat. Tras. (R$)", "Fat. Total (R$)"]
    cab(ws2, head2, 1)
    for i, (_, row) in enumerate(mensal.iterrows()):
        r = 2 + i
        vals = [row["lbl"], row["passagens"], row["refil_diant"], row["aproveitamento"],
                row["refil_tras"], row["aproveitamento_tras"], row["total_diant"],
                row["total_tras"], row["total_geral"]]
        for j, v in enumerate(vals, 1):
            c = ws2.cell(r, j, _v(v)); c.font = Font(name="Arial", size=10); c.border = BD
            if j in (2, 3, 5):
                c.number_format = "#,##0"; c.alignment = Alignment(horizontal="center")
            elif j in (4, 6):
                c.number_format = "0.0%"; c.alignment = Alignment(horizontal="center")
            elif j in (7, 8, 9):
                c.number_format = "R$ #,##0.00"
    for j, w in enumerate([12, 11, 12, 12, 11, 12, 14, 14, 14], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gerar_excel_semanal(g, ov, semana_lbl, ordenar):
    """Exporta o relatório semanal por gerente em Excel formatado."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    NAVY, WHITE = "FF1F3864", "FFFFFFFF"
    thin = Side(style="thin", color="FFD9D9D9")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _v(x):
        return None if (isinstance(x, float) and pd.isna(x)) else x

    wb = Workbook(); ws = wb.active; ws.title = "Relatório Semanal"
    ws.cell(1, 1, f"Resultado DAHRUJ {semana_lbl}").font = \
        Font(name="Arial", bold=True, size=13, color=NAVY)
    ws.cell(2, 1, f"Ordenado por {ordenar} · Gerado em "
                  f"{dt.datetime.now().strftime('%d/%m/%Y %H:%M')}").font = \
        Font(name="Arial", size=9, color="FF666666")
    head = ["Seq", "Gerente", "Marca", "Loja", "Passagens", "Refil Diant.", "% Aprov",
            "Total Diant.", "Refil Tras.", "% Aprov", "Total Tras.", "Total Geral", "Part %"]
    H = 4
    for j, h in enumerate(head, 1):
        c = ws.cell(H, j, h); c.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=NAVY); c.border = BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, r) in enumerate(g.iterrows()):
        rr = H + 1 + i
        vals = [i + 1, r["gerente"], r["marca"], r["loja"], r["passagens"], r["refil_diant"],
                r["aprov_d"], r["total_diant"], r["refil_tras"], r["aprov_t"],
                r["total_tras"], r["total_geral"], r["part"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(rr, j, _v(v)); c.font = Font(name="Arial", size=10); c.border = BD
            if j in (1, 5, 6, 7, 9, 10, 13):
                c.alignment = Alignment(horizontal="center")
            if j in (5, 6, 9):
                c.number_format = "#,##0"
            elif j in (7, 10, 13):
                c.number_format = "0.0%"
            elif j in (8, 11, 12):
                c.number_format = "R$ #,##0.00"
        if i % 2 == 1:
            for j in range(1, len(head) + 1):
                ws.cell(rr, j).fill = PatternFill("solid", fgColor="FFF4F6FA")
    tr = H + 1 + len(g)
    ws.cell(tr, 2, "TOTAL").font = Font(name="Arial", bold=True, size=10)
    tvals = {5: _v(ov["passagens"]), 6: _v(ov["refil_diant"]), 7: _v(ov["aprov_d"]),
             8: _v(ov["total_diant"]), 9: _v(ov["refil_tras"]), 10: _v(ov["aprov_t"]),
             11: _v(ov["total_tras"]), 12: _v(ov["total_geral"]), 13: 1.0}
    for j, v in tvals.items():
        c = ws.cell(tr, j, v); c.font = Font(name="Arial", bold=True, size=10); c.border = BD
        if j in (5, 6, 9):
            c.number_format = "#,##0"; c.alignment = Alignment(horizontal="center")
        elif j in (7, 10, 13):
            c.number_format = "0.0%"; c.alignment = Alignment(horizontal="center")
        elif j in (8, 11, 12):
            c.number_format = "R$ #,##0.00"
    for j, w in enumerate([5, 16, 8, 16, 11, 12, 9, 13, 11, 9, 13, 14, 8], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"A{H + 1}"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def gerar_excel_consultor(g, ov, mes_lbl, ordenar):
    """Exporta o relatório mensal por consultor em Excel formatado."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    NAVY, WHITE = "FF1F3864", "FFFFFFFF"
    thin = Side(style="thin", color="FFD9D9D9")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _v(x):
        return None if (isinstance(x, float) and pd.isna(x)) else x

    wb = Workbook(); ws = wb.active; ws.title = "Relatório Consultor"
    ws.cell(1, 1, f"Resultado por Consultor DAHRUJ {mes_lbl}").font = \
        Font(name="Arial", bold=True, size=13, color=NAVY)
    ws.cell(2, 1, f"Ordenado por {ordenar} · Gerado em "
                  f"{dt.datetime.now().strftime('%d/%m/%Y %H:%M')}").font = \
        Font(name="Arial", size=9, color="FF666666")
    head = ["Seq", "Consultor", "Unidade", "Passagens", "Refil Diant.", "% Aprov",
            "Total Diant.", "Refil Tras.", "% Aprov", "Total Tras.", "Total Geral",
            "Part %"]
    COL_INT, COL_PCT, COL_MONEY = (4, 5, 8), (6, 9, 12), (7, 10, 11)
    H = 4
    for j, h in enumerate(head, 1):
        c = ws.cell(H, j, h); c.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=NAVY); c.border = BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, r) in enumerate(g.iterrows()):
        rr = H + 1 + i
        vals = [i + 1, r["consultor"], r["unidade"], r["passagens"], r["refil_diant"],
                r["aprov_d"], r["total_diant"], r["refil_tras"], r["aprov_t"],
                r["total_tras"], r["total_geral"], r["part"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(rr, j, _v(v)); c.font = Font(name="Arial", size=10); c.border = BD
            if j == 1 or j in COL_INT or j in COL_PCT:
                c.alignment = Alignment(horizontal="center")
            if j in COL_INT:
                c.number_format = "#,##0"
            elif j in COL_PCT:
                c.number_format = "0.0%"
            elif j in COL_MONEY:
                c.number_format = "R$ #,##0.00"
        if i % 2 == 1:
            for j in range(1, len(head) + 1):
                ws.cell(rr, j).fill = PatternFill("solid", fgColor="FFF4F6FA")
    tr = H + 1 + len(g)
    ws.cell(tr, 2, "TOTAL").font = Font(name="Arial", bold=True, size=10)
    tvals = {4: _v(ov["passagens"]), 5: _v(ov["refil_diant"]), 6: _v(ov["aprov_d"]),
             7: _v(ov["total_diant"]), 8: _v(ov["refil_tras"]), 9: _v(ov["aprov_t"]),
             10: _v(ov["total_tras"]), 11: _v(ov["total_geral"]), 12: 1.0}
    for j, v in tvals.items():
        c = ws.cell(tr, j, v); c.font = Font(name="Arial", bold=True, size=10); c.border = BD
        if j in COL_INT:
            c.number_format = "#,##0"; c.alignment = Alignment(horizontal="center")
        elif j in COL_PCT:
            c.number_format = "0.0%"; c.alignment = Alignment(horizontal="center")
        elif j in COL_MONEY:
            c.number_format = "R$ #,##0.00"
    for j, w in enumerate([5, 22, 20, 11, 12, 9, 13, 11, 9, 13, 14, 8], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"A{H + 1}"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


